import os
import json
import shutil
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for
import config
from modules import photo_manager, ai_analyzer, excel_builder, lv_parser

app = Flask(__name__)
app.jinja_env.globals["basis_pfad"] = config.BASIS_PFAD

# ──────────────────────────────────────────────
# Hilfsfunktionen
# ──────────────────────────────────────────────

def liste_projekte() -> list[str]:
    if not os.path.exists(config.BASIS_PFAD):
        return []
    return [p for p in os.listdir(config.BASIS_PFAD)
            if os.path.isdir(os.path.join(config.BASIS_PFAD, p))]


def projekt_pfad(name: str) -> str:
    return os.path.join(config.BASIS_PFAD, name)


def lade_projekt_info(name: str) -> dict:
    info_datei = os.path.join(projekt_pfad(name), "projekt_info.json")
    if os.path.exists(info_datei):
        with open(info_datei) as f:
            return json.load(f)
    return {"auftraggeber": "", "baustelle": name, "bearbeiter": ""}


def speichere_projekt_info(name: str, info: dict):
    with open(os.path.join(projekt_pfad(name), "projekt_info.json"), "w") as f:
        json.dump(info, f, indent=2, ensure_ascii=False)


def lade_lv_fuer_ki(name: str) -> str:
    """Lädt LV-Positionen aus JSON (gecacht) und gibt kompakten KI-Text zurück."""
    p_pfad = projekt_pfad(name)
    positionen = lv_parser.lade_lv_json(p_pfad)
    if positionen:
        return lv_parser.lv_fuer_ki(positionen)
    return ""


def naechste_bericht_nr(name: str) -> int:
    zaehler_datei = os.path.join(config.BASIS_PFAD, "bericht_zaehler.json")
    try:
        with open(zaehler_datei) as f:
            data = json.load(f)
    except Exception:
        data = {}
    nr = data.get(name, 0) + 1
    data[name] = nr
    with open(zaehler_datei, "w") as f:
        json.dump(data, f, indent=2)
    return nr


def lade_status(name: str) -> dict:
    sf = os.path.join(projekt_pfad(name), "status.json")
    try:
        with open(sf) as f:
            return json.load(f)
    except Exception:
        return {}


def speichere_status(name: str, status: dict):
    sf = os.path.join(projekt_pfad(name), "status.json")
    with open(sf, "w") as f:
        json.dump(status, f, indent=2, ensure_ascii=False)


def finde_vorlage(name: str, typ: str) -> str | None:
    vorlagen_pfad = os.path.join(projekt_pfad(name),
                                 "1.4 Berichte", "1.4.1 Tagesberichte", "Vorlagen")
    if not os.path.exists(vorlagen_pfad):
        return None
    suchbegriff = "tages" if typ == "bautagesbericht" else "regie"
    for f in os.listdir(vorlagen_pfad):
        if f.lower().endswith((".xlsx", ".xlsm")) and not f.startswith("~$"):
            if suchbegriff in f.lower():
                return os.path.join(vorlagen_pfad, f)
    # Fallback: erste Excel-Datei
    for f in os.listdir(vorlagen_pfad):
        if f.lower().endswith((".xlsx", ".xlsm")) and not f.startswith("~$"):
            return os.path.join(vorlagen_pfad, f)
    return None


# ──────────────────────────────────────────────
# Routen
# ──────────────────────────────────────────────

@app.route("/")
def index():
    projekte = liste_projekte()
    return render_template("index.html", projekte=projekte)


@app.route("/fotos/<name>/<datum>")
def fotos_tag(name: str, datum: str):
    """Zeigt alle Fotos eines Tages als Vorschau."""
    fotos_pro_tag = photo_manager.lade_fotos(projekt_pfad(name))
    bilder = fotos_pro_tag.get(datum, [])
    dt = datetime.strptime(datum, "%Y-%m-%d")
    return render_template("fotos.html",
                           projekt=name,
                           datum=datum,
                           datum_anzeige=dt.strftime("%d.%m.%Y"),
                           bilder=bilder)


@app.route("/projekt/<name>")
def projekt(name: str):
    if name not in liste_projekte():
        return "Projekt nicht gefunden", 404

    fotos_pro_tag = photo_manager.lade_fotos(projekt_pfad(name))
    status = lade_status(name)

    tage = []
    for datum in sorted(fotos_pro_tag.keys(), reverse=True):
        dt = datetime.strptime(datum, "%Y-%m-%d")
        tage.append({
            "datum_key": datum,
            "datum_anzeige": dt.strftime("%d.%m.%Y"),
            "wochentag": ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"][dt.weekday()],
            "anzahl_fotos": len(fotos_pro_tag[datum]),
            "status": status.get(datum, "offen")
        })

    info = lade_projekt_info(name)
    return render_template("projekt.html", projekt=name, tage=tage, info=info)


@app.route("/projekt/<name>/einstellungen", methods=["GET", "POST"])
def projekt_einstellungen(name: str):
    if request.method == "POST":
        info = {
            "auftraggeber": request.form.get("auftraggeber", ""),
            "baustelle": request.form.get("baustelle", name),
            "bearbeiter": request.form.get("bearbeiter", "")
        }
        speichere_projekt_info(name, info)
        return redirect(url_for("projekt", name=name))
    info = lade_projekt_info(name)
    return render_template("einstellungen.html", projekt=name, info=info)


@app.route("/analysiere/<name>/<datum>", methods=["POST"])
def analysiere(name: str, datum: str):
    """KI-Analyse starten und Entwurf zurückgeben."""
    report_type = request.form.get("report_type", "bautagesbericht")

    fotos_pro_tag = photo_manager.lade_fotos(projekt_pfad(name))
    bilder = fotos_pro_tag.get(datum, [])

    if not bilder:
        return jsonify({"fehler": "Keine Fotos für dieses Datum gefunden."}), 400

    lv_text = lade_lv_fuer_ki(name)
    stichworte = request.form.get("stichworte", "").strip()

    try:
        ki_daten = ai_analyzer.analyze(bilder, lv_text, report_type, stichworte)
    except ValueError as e:
        return jsonify({"fehler": str(e)}), 400
    except Exception as e:
        return jsonify({"fehler": f"KI-Fehler: {str(e)}"}), 500

    # Entwurf zwischenspeichern
    entwurf_id = f"{name}__{datum}__{report_type}"
    entwurf_datei = os.path.join(config.BASIS_PFAD, f"entwurf_{entwurf_id}.json")
    with open(entwurf_datei, "w") as f:
        json.dump({
            "projekt": name,
            "datum": datum,
            "report_type": report_type,
            "bilder": bilder,
            "ki_daten": ki_daten
        }, f, indent=2, ensure_ascii=False)

    return redirect(url_for("entwurf", entwurf_id=entwurf_id))


@app.route("/entwurf/<entwurf_id>")
def entwurf(entwurf_id: str):
    entwurf_datei = os.path.join(config.BASIS_PFAD, f"entwurf_{entwurf_id}.json")
    if not os.path.exists(entwurf_datei):
        return "Entwurf nicht gefunden", 404
    with open(entwurf_datei) as f:
        daten = json.load(f)
    return render_template("entwurf.html", entwurf_id=entwurf_id, **daten)


@app.route("/bestaetigen/<entwurf_id>", methods=["POST"])
def bestaetigen(entwurf_id: str):
    """Korrigierten Entwurf speichern und Excel generieren."""
    entwurf_datei = os.path.join(config.BASIS_PFAD, f"entwurf_{entwurf_id}.json")
    if not os.path.exists(entwurf_datei):
        return jsonify({"fehler": "Entwurf nicht gefunden"}), 404

    with open(entwurf_datei) as f:
        entwurf_daten = json.load(f)

    name = entwurf_daten["projekt"]
    datum_key = entwurf_daten["datum"]
    report_type = entwurf_daten["report_type"]

    # Korrekturen aus Formular übernehmen
    ki_daten = entwurf_daten["ki_daten"]
    ki_daten["wetter_vormittag"] = request.form.get("wetter_vormittag", ki_daten.get("wetter_vormittag", ""))
    ki_daten["wetter_nachmittag"] = request.form.get("wetter_nachmittag", ki_daten.get("wetter_nachmittag", ""))
    ki_daten["temp_min"] = request.form.get("temp_min", ki_daten.get("temp_min", 0))
    ki_daten["temp_max"] = request.form.get("temp_max", ki_daten.get("temp_max", 0))
    ki_daten["personal_aufsicht"] = int(request.form.get("personal_aufsicht", 0))
    ki_daten["personal_facharbeiter"] = int(request.form.get("personal_facharbeiter", 0))
    ki_daten["personal_maschinist"] = int(request.form.get("personal_maschinist", 0))

    # Mehrzeilige Felder
    for feld in ("beschreibung_arbeiten", "geraete_liste", "material_liste",
                 "lv_positionen", "sonstiges"):
        rohtext = request.form.get(feld, "")
        ki_daten[feld] = [z.strip() for z in rohtext.splitlines() if z.strip()]

    # Regiebericht-spezifisch
    if report_type == "regiebericht":
        ki_daten["grund_regie"] = request.form.get("grund_regie", ki_daten.get("grund_regie", ""))

    # Datum formatieren
    dt = datetime.strptime(datum_key, "%Y-%m-%d")
    datum_fmt = dt.strftime("%d.%m.%Y")

    # Vorlage und Ausgabepfad
    vorlage = finde_vorlage(name, report_type)
    if not vorlage:
        return jsonify({"fehler": "Keine Excel-Vorlage gefunden"}), 400

    ext = os.path.splitext(vorlage)[1]
    ausgabe_pfad = excel_builder.get_output_path(projekt_pfad(name), datum_fmt, report_type, ext)
    bericht_nr = naechste_bericht_nr(name)
    info = lade_projekt_info(name)

    if report_type == "bautagesbericht":
        excel_builder.bautagesbericht(vorlage, ausgabe_pfad, ki_daten, datum_fmt, bericht_nr, info)
    else:
        excel_builder.regiebericht(vorlage, ausgabe_pfad, ki_daten, datum_fmt, bericht_nr, info)

    # Status aktualisieren
    status = lade_status(name)
    status[datum_key] = "fertig"
    speichere_status(name, status)

    # Entwurf-Datei aufräumen
    os.remove(entwurf_datei)

    # Bilder in Ausgabe-Ordner kopieren
    bild_ziel = os.path.join(os.path.dirname(ausgabe_pfad), "Bilder")
    os.makedirs(bild_ziel, exist_ok=True)
    for bild in entwurf_daten.get("bilder", []):
        try:
            shutil.copy2(bild, bild_ziel)
        except Exception:
            pass

    return redirect(url_for("download", name=name, dateiname=os.path.basename(ausgabe_pfad),
                             unterordner=os.path.relpath(os.path.dirname(ausgabe_pfad),
                                                          projekt_pfad(name))))


@app.route("/download/<name>/<path:unterordner>/<dateiname>")
def download(name: str, unterordner: str, dateiname: str):
    pfad = os.path.join(projekt_pfad(name), unterordner, dateiname)
    if not os.path.exists(pfad):
        return "Datei nicht gefunden", 404
    return send_file(pfad, as_attachment=True, download_name=dateiname)


@app.route("/debug_vorlage/<name>")
def debug_vorlage(name: str):
    """Zeigt Struktur der Excel-Vorlage – hilft bei Zell-Mapping."""
    import openpyxl
    vorlage = finde_vorlage(name, "bautagesbericht")
    if not vorlage:
        return jsonify({"fehler": "Keine Vorlage gefunden", "suchpfad": os.path.join(projekt_pfad(name), "1.4 Berichte", "1.4.1 Tagesberichte", "Vorlagen")})

    wb = openpyxl.load_workbook(vorlage, keep_vba=True)
    ergebnis = {"vorlage_pfad": vorlage, "tabs": {}}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        merged = [str(r) for r in ws.merged_cells.ranges]
        # Erste 50 Zellen mit Inhalt
        zellen = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and str(cell.value).strip():
                    zellen[cell.coordinate] = str(cell.value)[:60]
                if len(zellen) >= 50:
                    break
            if len(zellen) >= 50:
                break
        ergebnis["tabs"][sheet_name] = {
            "merged_cells": merged[:30],
            "zellen_mit_inhalt": zellen
        }

    return jsonify(ergebnis)


@app.route("/lv/<name>")
def lv_positionen(name: str):
    """Gibt alle LV-Positionen als JSON zurück (für Dropdown im Entwurf)."""
    positionen = lv_parser.lade_lv_json(projekt_pfad(name))
    return jsonify(positionen)


@app.route("/foto/<path:pfad>")
def foto(pfad: str):
    """Liefert ein Foto als Vorschau aus."""
    voller_pfad = os.path.join(config.BASIS_PFAD, pfad)
    if os.path.exists(voller_pfad):
        return send_file(voller_pfad)
    return "Foto nicht gefunden", 404


if __name__ == "__main__":
    os.makedirs(config.BASIS_PFAD, exist_ok=True)
    app.run(host="0.0.0.0", port=5000, debug=True)
